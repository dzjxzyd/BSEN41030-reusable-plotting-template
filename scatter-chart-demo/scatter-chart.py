#!/usr/bin/env python3
"""Create a muted, publication-style scatter chart from an Excel data table.

Expected workbook layout (first worksheet):

    | Sample_ID | Architecture | t-SNE_1 | t-SNE_2 | ... |
    | MLP-001   | MLP           | -6.25   | -5.64   | ... |

The label, x-coordinate, and y-coordinate column names are supplied by the
caller. Every data row contributes one point, so groups may have different
numbers of points.
"""

from __future__ import annotations

import argparse
import math
import sys
from collections import OrderedDict
from pathlib import Path
from typing import Any, Mapping


DEFAULT_X_LABEL = "t-SNE_1"
DEFAULT_Y_LABEL = "t-SNE_2"


class ExcelFormatError(ValueError):
    """Raised when the workbook does not match the expected table layout."""


def _is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _trim_table(rows: list[list[Any]]) -> list[list[Any]]:
    """Remove completely empty trailing rows and columns."""
    while rows and all(_is_blank(value) for value in rows[-1]):
        rows.pop()
    if not rows:
        return rows

    width = max(len(row) for row in rows)
    for row in rows:
        row.extend([None] * (width - len(row)))

    while width and all(_is_blank(row[width - 1]) for row in rows):
        width -= 1
    return [row[:width] for row in rows]


def _resolve_column(headers: list[str], requested: str, role: str) -> int:
    requested = requested.strip() if isinstance(requested, str) else ""
    if not requested:
        raise ExcelFormatError(f"The {role} column name cannot be empty.")

    matches = [index for index, header in enumerate(headers) if header == requested]
    if not matches:
        available = ", ".join(repr(header) for header in headers if header)
        raise ExcelFormatError(
            f"The {role} column {requested!r} was not found. Available columns: {available}."
        )
    if len(matches) > 1:
        raise ExcelFormatError(f"The {role} column {requested!r} appears more than once.")
    return matches[0]


def _to_coordinate(value: Any, column_name: str, row_number: int) -> float:
    if _is_blank(value):
        raise ExcelFormatError(
            f"Invalid coordinate at Excel row {row_number}, column {column_name!r}: value is empty."
        )
    if isinstance(value, bool):
        raise ExcelFormatError(
            f"Invalid coordinate at Excel row {row_number}, column {column_name!r}: {value!r} is not numeric."
        )
    try:
        number = float(value)
    except (TypeError, ValueError):
        raise ExcelFormatError(
            f"Invalid coordinate at Excel row {row_number}, column {column_name!r}: {value!r} is not numeric."
        ) from None
    if not math.isfinite(number):
        raise ExcelFormatError(
            f"Invalid coordinate at Excel row {row_number}, column {column_name!r}: {value!r} is not finite."
        )
    return number


def read_excel_file(
    path: str | Path,
    label_column: str,
    x_column: str,
    y_column: str,
) -> OrderedDict[str, list[tuple[float, float]]]:
    """Read the first worksheet into an ordered label -> point mapping."""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("Missing dependency: install openpyxl with `python -m pip install openpyxl`.") from exc

    workbook_path = Path(path).expanduser()
    if not workbook_path.exists():
        raise FileNotFoundError(f"Input Excel file does not exist: {workbook_path}")
    if workbook_path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ExcelFormatError("Please provide an .xlsx or .xlsm Excel file.")

    try:
        workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    except Exception as exc:
        raise ExcelFormatError(f"Could not read workbook {workbook_path}: {exc}") from exc

    try:
        if not workbook.worksheets:
            raise ExcelFormatError("The workbook does not contain a worksheet.")
        worksheet = workbook.worksheets[0]
        rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
    finally:
        workbook.close()
    rows = _trim_table(rows)
    if len(rows) < 2 or not rows[0]:
        raise ExcelFormatError("The first worksheet must contain a header row and at least one data row.")

    headers = [str(value).strip() if not _is_blank(value) else "" for value in rows[0]]
    non_blank_headers = [header for header in headers if header]
    duplicates = sorted({header for header in non_blank_headers if non_blank_headers.count(header) > 1})
    if duplicates:
        names = ", ".join(repr(header) for header in duplicates)
        raise ExcelFormatError(f"Duplicate column headers found: {names}.")

    label_index = _resolve_column(headers, label_column, "label")
    x_index = _resolve_column(headers, x_column, "x-coordinate")
    y_index = _resolve_column(headers, y_column, "y-coordinate")

    groups: OrderedDict[str, list[tuple[float, float]]] = OrderedDict()
    for row_number, row in enumerate(rows[1:], start=2):
        label_value = row[label_index]
        x_value = row[x_index]
        y_value = row[y_index]
        if _is_blank(label_value) and _is_blank(x_value) and _is_blank(y_value):
            continue
        if _is_blank(label_value):
            raise ExcelFormatError(f"Label is empty at Excel row {row_number}.")

        label = str(label_value).strip()
        x_value = _to_coordinate(x_value, headers[x_index], row_number)
        y_value = _to_coordinate(y_value, headers[y_index], row_number)
        groups.setdefault(label, []).append((x_value, y_value))

    if not groups:
        raise ExcelFormatError("No data rows with labels and coordinates were found.")
    return groups


_MUTED_PALETTE = (
    "#2F8991",
    "#DC7440",
    "#5B72B5",
    "#91A64F",
    "#A66A9B",
    "#BE9347",
    "#B85D68",
    "#5C89A1",
    "#729A73",
    "#9B75A7",
    "#8B765E",
    "#688E8C",
)


def _palette(size: int) -> list[Any]:
    """Return deterministic muted colors, expanding gently for many groups."""
    if size <= len(_MUTED_PALETTE):
        return list(_MUTED_PALETTE[:size])

    from matplotlib.colors import hsv_to_rgb

    colors = list(_MUTED_PALETTE)
    for index in range(len(_MUTED_PALETTE), size):
        hue = (index / size + 0.04) % 1.0
        colors.append(tuple(hsv_to_rgb((hue, 0.34, 0.76))))
    return colors


def make_plot(
    groups: Mapping[str, list[tuple[float, float]]],
    output_path: str | Path,
    x_label: str,
    y_label: str,
) -> Path:
    """Render and save a scatter chart using caller-provided axis titles."""
    if not groups:
        raise ExcelFormatError("Cannot render a scatter chart without any labeled points.")

    try:
        import matplotlib.pyplot as plt
    except ImportError as exc:
        raise RuntimeError(
            "Missing dependency: install matplotlib with `python -m pip install matplotlib`."
        ) from exc

    output_path = Path(output_path).expanduser()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    total_points = sum(len(points) for points in groups.values())
    point_size = 30 if total_points <= 1000 else 18 if total_points <= 5000 else 10
    figure_width = max(7.2, min(14.0, 8.0 + 0.55 * len(groups)))

    plt.rcParams.update(
        {
            "font.family": "DejaVu Sans",
            "font.size": 15,
            "axes.labelsize": 15,
            "axes.labelweight": "normal",
            "xtick.labelsize": 15,
            "ytick.labelsize": 15,
            "axes.linewidth": 1.5,
        }
    )
    fig, ax = plt.subplots(figsize=(figure_width, 5.8))
    fig.subplots_adjust(left=0.16, bottom=0.18, right=0.76 if len(groups) > 1 else 0.95, top=0.93)

    for color, (label, points) in zip(_palette(len(groups)), groups.items()):
        x_values = [point[0] for point in points]
        y_values = [point[1] for point in points]
        ax.scatter(
            x_values,
            y_values,
            s=point_size,
            color=color,
            alpha=0.82,
            edgecolors="white",
            linewidths=0.45,
            label=label,
            zorder=3,
        )

    ax.set_ylabel(y_label, fontweight="normal")
    ax.set_xlabel(x_label, fontweight="normal")
    for spine in ax.spines.values():
        spine.set_visible(True)
        spine.set_linewidth(1.5)
        spine.set_color("#1F2933")
    ax.grid(False)
    ax.tick_params(axis="both", length=5, width=1.3)
    ax.margins(0.06)
    if len(groups) > 1:
        ax.legend(
            frameon=False,
            loc="upper left",
            bbox_to_anchor=(1.01, 1.0),
            borderaxespad=0,
            markerscale=1.15,
        )

    fig.savefig(output_path, dpi=600, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return output_path


def run(
    input_path: str | Path,
    output_path: str | Path | None = None,
    label_column: str | None = None,
    x_column: str | None = None,
    y_column: str | None = None,
    x_label: str | None = None,
    y_label: str | None = None,
) -> Path:
    """Read the workbook and render a chart, prompting for omitted values."""
    input_path = Path(input_path).expanduser()
    if label_column is None:
        label_column = input("Label column name: ").strip()
    if x_column is None:
        x_column = input("X-coordinate column name: ").strip()
    if y_column is None:
        y_column = input("Y-coordinate column name: ").strip()
    if x_label is None:
        entered = input(f"Enter the x-axis title [{DEFAULT_X_LABEL}]: ").strip()
        x_label = entered or DEFAULT_X_LABEL
    if y_label is None:
        entered = input(f"Enter the y-axis title [{DEFAULT_Y_LABEL}]: ").strip()
        y_label = entered or DEFAULT_Y_LABEL

    groups = read_excel_file(input_path, label_column, x_column, y_column)
    output_path = Path(output_path).expanduser() if output_path else input_path.with_name(f"{input_path.stem}_tsne_scatter.png")
    return make_plot(groups, output_path, x_label, y_label)


def main(
    excel_file: str | Path | None = None,
    output_path: str | Path | None = None,
    label_column: str | None = None,
    x_column: str | None = None,
    y_column: str | None = None,
    x_label: str | None = None,
    y_label: str | None = None,
) -> int:
    """Run the scatter chart from a notebook or command line."""
    if excel_file is None:
        entered_filename = input("Excel filename (path or filename in this folder): ").strip()
        if not entered_filename:
            print("Error: no Excel filename was provided.", file=sys.stderr)
            return 1
        excel_file = entered_filename

    try:
        output = run(excel_file, output_path, label_column, x_column, y_column, x_label, y_label)
    except (ExcelFormatError, FileNotFoundError, RuntimeError) as exc:
        print(f"Error: {exc}", file=sys.stderr)
        return 1

    print(f"PNG saved to: {output}")
    return 0


def cli_main() -> int:
    """Parse command-line options, then delegate to :func:`main`."""
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "excel_file",
        type=Path,
        nargs="?",
        help="Input .xlsx or .xlsm file; if omitted, the script asks for it",
    )
    parser.add_argument("-o", "--output", type=Path, help="PNG output path")
    parser.add_argument("--label-column", help="Column used to group points")
    parser.add_argument("--x-column", help="Column containing x coordinates")
    parser.add_argument("--y-column", help="Column containing y coordinates")
    parser.add_argument("--x-label", help="Title displayed on the x axis")
    parser.add_argument("--y-label", help="Title displayed on the y axis")
    args = parser.parse_args()
    return main(
        args.excel_file,
        args.output,
        args.label_column,
        args.x_column,
        args.y_column,
        args.x_label,
        args.y_label,
    )


if __name__ == "__main__":
    raise SystemExit(cli_main())
