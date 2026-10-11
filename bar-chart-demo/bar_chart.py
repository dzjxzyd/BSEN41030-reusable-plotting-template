#!/usr/bin/env python3
"""Create a publication-style bar chart from a replicate-formatted Excel file.

Expected workbook layout (first worksheet):

    |              | Replicate 1 | Replicate 2 | Replicate 3 |
    | Sample 1     | 10.1        | 10.4        | 10.2        |
    | Sample 2     | 12.8        | 12.5        | 12.7        |

Usage:
    python bar_chart.py "my_dataset.xlsx"

Dependencies:
    openpyxl, matplotlib, scipy
"""

from __future__ import annotations

import argparse
import math
import sys
from collections import OrderedDict
from pathlib import Path
from typing import Any, Iterable


class ExcelFormatError(ValueError):
    """Raised when the workbook does not match the expected tabular layout."""


def _is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _trim_table(rows: list[list[Any]]) -> list[list[Any]]:
    """Remove completely empty trailing rows and columns."""
    while rows and all(_is_blank(value) for value in rows[-1]):
        rows.pop()
    if not rows:
        return rows

    width = max(len(row) for row in rows)
    for row in rows:
        row.extend([None] * (width - len(row)))

    while width and all(_is_blank(row[width - 1]) for row in rows):
        width -= 1
    return [row[:width] for row in rows]


def _to_number(value: Any, sample: str, replicate: str, row_number: int, column_number: int) -> float | None:
    if _is_blank(value):
        return None
    if isinstance(value, bool):
        raise ExcelFormatError(
            f"Non-numeric value at Excel cell row {row_number}, column {column_number} "
            f"({sample!r}, {replicate!r}): {value!r}"
        )
    try:
        number = float(value)
    except (TypeError, ValueError):
        raise ExcelFormatError(
            f"Non-numeric value at Excel cell row {row_number}, column {column_number} "
            f"({sample!r}, {replicate!r}): {value!r}"
        ) from None
    if not math.isfinite(number):
        raise ExcelFormatError(
            f"Non-finite value at Excel cell row {row_number}, column {column_number} "
            f"({sample!r}, {replicate!r}): {value!r}"
        )
    return number


def read_excel_file(path: str | Path) -> OrderedDict[str, list[float]]:
    """Read the first worksheet into an ordered sample -> replicate values map."""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("Missing dependency: install openpyxl with `python -m pip install openpyxl`.") from exc

    workbook_path = Path(path).expanduser()
    if not workbook_path.exists():
        raise FileNotFoundError(f"Input Excel file does not exist: {workbook_path}")
    if workbook_path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ExcelFormatError("Please provide an .xlsx or .xlsm Excel file.")

    try:
        workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    except Exception as exc:
        raise ExcelFormatError(f"Could not read workbook {workbook_path}: {exc}") from exc

    try:
        if not workbook.worksheets:
            raise ExcelFormatError("The workbook does not contain a worksheet.")
        worksheet = workbook.worksheets[0]
        rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
    finally:
        workbook.close()
    rows = _trim_table(rows)
    if len(rows) < 2 or len(rows[0]) < 2:
        raise ExcelFormatError(
            "The worksheet must have a header row, a sample-name column, and at least one replicate column."
        )

    replicate_columns: list[tuple[int, str]] = []
    for column_index, value in enumerate(rows[0][1:], start=1):
        if _is_blank(value):
            # Empty trailing columns were already removed. An empty header in the
            # middle is ambiguous and should not be silently interpreted.
            if any(not _is_blank(row[column_index]) for row in rows[1:]):
                raise ExcelFormatError(
                    f"Replicate header is empty in Excel column {column_index + 1}, but that column contains data."
                )
            continue
        replicate_columns.append((column_index, str(value).strip()))

    if not replicate_columns:
        raise ExcelFormatError("No replicate headers were found in the first row.")

    samples: OrderedDict[str, list[float]] = OrderedDict()
    for row_index, row in enumerate(rows[1:], start=2):
        sample_cell = row[0]
        if _is_blank(sample_cell):
            if any(not _is_blank(row[column_index]) for column_index, _ in replicate_columns):
                raise ExcelFormatError(f"Sample name is empty in Excel row {row_index}, but that row contains data.")
            continue

        sample = str(sample_cell).strip()
        if sample in samples:
            raise ExcelFormatError(f"Duplicate sample name found: {sample!r}")

        values: list[float] = []
        for column_index, replicate in replicate_columns:
            value = _to_number(row[column_index], sample, replicate, row_index, column_index + 1)
            if value is not None:
                values.append(value)
        if not values:
            raise ExcelFormatError(f"Sample {sample!r} has no numeric replicate values.")
        samples[sample] = values

    if len(samples) < 2:
        raise ExcelFormatError("At least two samples with numeric values are required for significance testing.")
    return samples


def _mean(values: Iterable[float]) -> float:
    values = list(values)
    return sum(values) / len(values)


def _sample_sd(values: Iterable[float]) -> float:
    values = list(values)
    if len(values) < 2:
        return 0.0
    average = _mean(values)
    return math.sqrt(sum((value - average) ** 2 for value in values) / (len(values) - 1))


def calculate_summary(samples: OrderedDict[str, list[float]]) -> list[dict[str, Any]]:
    """Calculate n, mean, and sample SD while preserving Excel sample order."""
    return [
        {
            "sample": sample,
            "n": len(values),
            "mean": _mean(values),
            "sd": _sample_sd(values),
        }
        for sample, values in samples.items()
    ]


def _load_statistics_dependencies():
    try:
        from scipy import stats
    except ImportError as exc:
        raise RuntimeError(
            "Missing dependency: install scipy with `python -m pip install scipy`."
        ) from exc
    return stats


def _tukey_p_value(q: float, number_of_groups: int, degrees_of_freedom: int, stats: Any) -> float:
    studentized_range = getattr(stats, "studentized_range", None)
    if studentized_range is not None:
        return float(studentized_range.sf(q, number_of_groups, degrees_of_freedom))

    try:
        from statsmodels.stats.libqsturng import psturng
    except ImportError as exc:
        raise RuntimeError(
            "This SciPy version lacks studentized_range. Install statsmodels too with "
            "`python -m pip install statsmodels`."
        ) from exc
    return float(psturng(q, number_of_groups, degrees_of_freedom)[0])


def calculate_statistics(samples: OrderedDict[str, list[float]]) -> tuple[dict[str, float | str], list[dict[str, Any]]]:
    """Run one-way ANOVA followed by Tukey-Kramer HSD pairwise comparisons."""
    stats = _load_statistics_dependencies()
    names = list(samples)
    groups = [samples[name] for name in names]
    number_of_groups = len(groups)
    total_n = sum(len(group) for group in groups)
    grand_mean = sum(sum(group) for group in groups) / total_n
    between_ss = sum(len(group) * (_mean(group) - grand_mean) ** 2 for group in groups)
    within_ss = sum(sum((value - _mean(group)) ** 2 for value in group) for group in groups)
    degrees_between = number_of_groups - 1
    degrees_within = total_n - number_of_groups

    if degrees_within <= 0:
        raise ExcelFormatError(
            "ANOVA requires at least one residual degree of freedom; add replicate values."
        )
    mean_square_between = between_ss / degrees_between
    mean_square_within = within_ss / degrees_within
    if mean_square_within == 0:
        if between_ss == 0:
            f_value = 0.0
            anova_p = 1.0
        else:
            f_value = math.inf
            anova_p = 0.0
    else:
        f_value = mean_square_between / mean_square_within
        anova_p = float(stats.f.sf(f_value, degrees_between, degrees_within))

    anova = {"anova_f": f_value, "anova_p": anova_p}
    tukey: list[dict[str, Any]] = []
    for left_index, left_name in enumerate(names):
        for right_name in names[left_index + 1 :]:
            left = samples[left_name]
            right = samples[right_name]
            difference = _mean(left) - _mean(right)
            standard_error = math.sqrt((mean_square_within / 2.0) * (1 / len(left) + 1 / len(right)))
            if standard_error:
                q_value = abs(difference) / standard_error
                p_adjusted = _tukey_p_value(q_value, number_of_groups, degrees_within, stats)
            elif difference == 0:
                # Zero variance and identical means: no evidence of a difference.
                q_value = 0.0
                p_adjusted = 1.0
            else:
                q_value = math.inf
                p_adjusted = 0.0
            tukey.append(
                {
                    "group1": left_name,
                    "group2": right_name,
                    "mean_diff": difference,
                    "q": q_value,
                    "p_adj": p_adjusted,
                    "reject": bool(p_adjusted < 0.05),
                }
            )
    return anova, tukey


def significance_letters(summary: list[dict[str, Any]], tukey: list[dict[str, Any]]) -> dict[str, str]:
    """Assign compact letters; groups sharing a letter are not significantly different.

    Uses the insert-and-absorb algorithm (Piepho, 2004), so every pair of groups
    that is not significantly different shares at least one letter, and every
    significantly different pair shares none. Letter "a" marks the highest mean.
    """
    names = [item["sample"] for item in sorted(summary, key=lambda item: item["mean"], reverse=True)]
    rank = {name: index for index, name in enumerate(names)}
    significant_pairs = [(row["group1"], row["group2"]) for row in tukey if row["reject"]]

    columns: list[set[str]] = [set(names)]
    for first, second in significant_pairs:
        updated: list[set[str]] = []
        for column in columns:
            if first in column and second in column:
                updated.extend([column - {first}, column - {second}])
            else:
                updated.append(column)
        # Absorb: drop empty, duplicate, and subset columns.
        unique: list[set[str]] = []
        for column in sorted(updated, key=len, reverse=True):
            if column and not any(column <= kept for kept in unique):
                unique.append(column)
        columns = unique

    columns.sort(key=lambda column: sorted(rank[name] for name in column))
    if len(columns) > 26:
        raise ExcelFormatError("Too many distinct significance groups to label with a-z.")
    assignments: dict[str, str] = {name: "" for name in names}
    for letter_index, column in enumerate(columns):
        for name in sorted(column, key=rank.get):
            assignments[name] += chr(ord("a") + letter_index)
    return assignments


def _configure_x_tick_layout(fig: Any, ax: Any, names: list[str]) -> None:
    """Rotate x tick labels only when their rendered bounds would overlap."""
    ax.set_xticks(range(len(names)), names)
    fig.canvas.draw()

    renderer = fig.canvas.get_renderer()
    labels = ax.get_xticklabels()
    bounds = [label.get_window_extent(renderer) for label in labels]
    minimum_gap = 4.0 * fig.dpi / 72.0
    labels_overlap = any(
        left.x1 + minimum_gap > right.x0
        for left, right in zip(bounds, bounds[1:])
    )

    if not labels_overlap:
        fig.subplots_adjust(bottom=0.22)
        return

    for label in labels:
        label.set_rotation(45)
        label.set_horizontalalignment("right")
        label.set_rotation_mode("anchor")

    fig.canvas.draw()
    renderer = fig.canvas.get_renderer()
    label_height = max(
        (label.get_window_extent(renderer).height for label in labels),
        default=0.0,
    )
    bottom_margin = min(0.48, max(0.26, label_height / fig.bbox.height + 0.13))
    fig.subplots_adjust(bottom=bottom_margin)


def make_plot(
    samples: OrderedDict[str, list[float]],
    summary: list[dict[str, Any]],
    letters: dict[str, str],
    x_label: str,
    y_label: str,
    output_path: str | Path,
) -> Path:
    """Render and save the publication-style PNG."""
    try:
        import matplotlib.pyplot as plt
    except ImportError as exc:
        raise RuntimeError(
            "Missing dependency: install matplotlib with `python -m pip install matplotlib`."
        ) from exc

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    names = [row["sample"] for row in summary]
    means = [row["mean"] for row in summary]
    sds = [row["sd"] for row in summary]
    x_values = list(range(len(names)))
    palette = "#214A55"

    plt.rcParams.update({
        "font.family": "DejaVu Sans",
        "font.size": 15,
        "axes.labelsize": 15,
        "axes.labelweight": "normal",
        "xtick.labelsize": 15,
        "ytick.labelsize": 15,
        "axes.linewidth": 1.5,
    })
    fig_width = max(7.2, min(14.0, 1.0 + 0.78 * len(names)))
    fig, ax = plt.subplots(figsize=(fig_width, 5.8))
    fig.subplots_adjust(left=0.22, bottom=0.22, right=0.95, top=0.93)
    bars = ax.bar(
        x_values,
        means,
        yerr=sds,
        width=0.68,
        color=palette,
        alpha=0.88,
        edgecolor="#1F2933",
        linewidth=0.8,
        capsize=5,
        error_kw={"elinewidth": 1.35, "capthick": 1.35, "ecolor": "#1F2933"},
        zorder=2,
    )

    top_values = [row["mean"] + row["sd"] for row in summary]
    data_min = min(min(values) for values in samples.values())
    data_max = max(max(values) for values in samples.values())
    top = max(top_values + [data_max])
    bottom = min(0.0, data_min)
    span = max(top - bottom, 1.0)
    label_offset = span * 0.035
    for bar, row in zip(bars, summary):
        ax.text(
            bar.get_x() + bar.get_width() / 2,
            row["mean"] + row["sd"] + label_offset,
            letters.get(row["sample"], "—"),
            ha="center",
            va="bottom",
            fontsize=12,
            fontweight="normal",
            color="#1F2933",
        )

    ax.set_ylabel(y_label, fontsize=17, fontweight="normal")
    ax.set_xlabel(x_label, fontsize=17, fontweight="normal")
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.spines["left"].set_linewidth(1.5)
    ax.spines["bottom"].set_linewidth(1.5)
    ax.grid(False)
    ax.tick_params(axis="both", length=5, width=1.3)
    _configure_x_tick_layout(fig, ax, names)
    ax.set_ylim(bottom - 0.04 * span, top + 0.16 * span)

    fig.savefig(output_path, dpi=600, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return output_path


def _default_y_label(input_path: Path) -> str:
    return input_path.stem.replace("-", " ").replace("_", " ").strip() or "Measurement"


def run(
    input_path: str | Path,
    output_path: str | Path | None = None,
    x_label: str | None = None,
    y_label: str | None = None,
) -> Path:
    """Analyze the workbook and render a chart, prompting for omitted labels."""
    input_path = Path(input_path).expanduser()
    samples = read_excel_file(input_path)
    summary = calculate_summary(samples)
    anova, tukey = calculate_statistics(samples)
    letters = significance_letters(summary, tukey)

    if x_label is None:
        entered = input("Enter the x-axis title [Sample]: ").strip()
        x_label = entered or "Sample"

    if y_label is None:
        default_label = _default_y_label(input_path)
        entered = input(f"Enter the y-axis title [{default_label}]: ").strip()
        y_label = entered or default_label

    output_path = Path(output_path).expanduser() if output_path else input_path.with_name(f"{input_path.stem}_barplot.png")
    make_plot(samples, summary, letters, x_label, y_label, output_path)
    print_statistics(summary, anova, tukey, letters)
    return output_path


def print_statistics(
    summary: list[dict[str, Any]],
    anova: dict[str, float | str],
    tukey: list[dict[str, Any]],
    letters: dict[str, str],
) -> None:
    """Print the ANOVA result, group summary, and Tukey HSD comparisons."""
    print(f"One-way ANOVA: F = {anova['anova_f']:.3f}, p = {anova['anova_p']:.3g}")
    print("\nGroup summary (bars show mean ± SD):")
    for row in summary:
        print(
            f"  {row['sample']}: n = {row['n']}, mean = {row['mean']:.3f}, "
            f"SD = {row['sd']:.3f}, letter = {letters[row['sample']]}"
        )
    print("\nTukey HSD (groups sharing a letter are not significantly different, alpha = 0.05):")
    for row in tukey:
        verdict = "significant" if row["reject"] else "not significant"
        print(f"  {row['group1']} vs {row['group2']}: p_adj = {row['p_adj']:.4g} ({verdict})")


def main(
    excel_file: str | Path | None = None,
    output_path: str | Path | None = None,
    x_label: str | None = None,
    y_label: str | None = None,
) -> int:
    """Run the analysis from a notebook or from the command-line wrapper.

    In a notebook:

        filename = None  # or: "sample-1-antioxidant.xlsx"
        main(filename, x_label=None, y_label=None)

    If ``excel_file`` is ``None``, the user is prompted for the filename.
    If either axis label is ``None``, the user is prompted for that title.
    """
    if excel_file is None:
        entered_filename = input("Excel filename (path or filename in this folder): ").strip()
        if not entered_filename:
            print("Error: no Excel filename was provided.", file=sys.stderr)
            return 1
        excel_file = entered_filename

    try:
        output_path = run(excel_file, output_path, x_label, y_label)
    except (ExcelFormatError, FileNotFoundError, RuntimeError) as exc:
        print(f"Error: {exc}", file=sys.stderr)
        return 1

    print(f"PNG saved to: {output_path}")
    return 0


def cli_main() -> int:
    """Parse command-line options, then delegate to the notebook-friendly main."""
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "excel_file",
        type=Path,
        nargs="?",
        help="Input .xlsx or .xlsm file; if omitted, the script asks for it",
    )
    parser.add_argument("-o", "--output", type=Path, help="PNG output path")
    parser.add_argument("--x-label", help="X-axis title; omit to be prompted before plotting")
    parser.add_argument("--y-label", help="Y-axis title; omit to be prompted before plotting")
    args = parser.parse_args()
    return main(args.excel_file, args.output, args.x_label, args.y_label)


if __name__ == "__main__":
    raise SystemExit(cli_main())
